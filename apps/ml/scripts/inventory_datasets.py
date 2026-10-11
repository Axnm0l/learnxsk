"""Inventory the read-only dataset directory without modifying it.

Phase 1 responsibility: produce a machine-readable description of every dataset
grouping found under the datasets root. This script never writes into the
datasets directory; it only reads, hashes, and samples.

Grouping rule
-------------
A *dataset* is one immediate subdirectory of ``--datasets-root``. Loose files
directly in the root (for example ``README.md``) are reported separately as
``ungrouped`` so they are never miscounted as datasets.

Usage (Windows PowerShell)::

    python apps\\ml\\scripts\\inventory_datasets.py `
        --datasets-root "C:\\Users\\anmol\\OneDrive\\Desktop\\hackathon\\datasets"

Outputs a JSON inventory under the validated training root
(``reports/data_inventory.json`` by default) and prints a short summary.

Standard library only, so it runs before any ML dependency is installed.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _common  # noqa: E402  (path bootstrap must precede import)

MAX_SAMPLE_CHARS = 400
MAX_SAMPLES_PER_FILE = 3
HASH_CHUNK = 1024 * 1024

TABULAR_EXT = {".jsonl", ".ndjson", ".json", ".csv", ".tsv"}
ARCHIVE_EXT = {".zip", ".tar", ".gz", ".tgz", ".bz2", ".xz", ".7z", ".rar", ".zst"}
BINARY_DOC_EXT = {".pdf", ".xlsx", ".xls", ".docx", ".pptx", ".png", ".jpg", ".jpeg"}

# Patterns used only to *flag* possible secrets. Matches are masked, never
# printed in full, per the task's safety rules.
SECRET_PATTERNS = (
    ("email", re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")),
    ("openai_key", re.compile(r"\bsk-[A-Za-z0-9_\-]{16,}\b")),
    ("github_token", re.compile(r"\bgh[pousr]_[A-Za-z0-9]{20,}\b")),
    ("aws_key", re.compile(r"\bAKIA[0-9A-Z]{16}\b")),
    ("private_key", re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----")),
    ("bearer", re.compile(r"(?i)\bbearer\s+[A-Za-z0-9._\-]{20,}")),
    ("password_assign", re.compile(r"(?i)\b(password|passwd|secret|api[_-]?key|token)\b\s*[:=]\s*\S{6,}")),
    ("slack_token", re.compile(r"\bxox[abprs]-[A-Za-z0-9\-]{10,}\b")),
)


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(HASH_CHUNK)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def mask_secrets(text: str) -> tuple[str, list[str]]:
    """Return (masked_text, sorted list of matched category names)."""
    hits: list[str] = []
    masked = text
    for label, pattern in SECRET_PATTERNS:
        def _sub(match: re.Match[str], _label: str = label) -> str:
            hits.append(_label)
            return f"<<REDACTED:{_label}>>"
        masked = pattern.sub(_sub, masked)
    return masked, sorted(set(hits))


def redact_preview(value: Any) -> tuple[Any, list[str]]:
    """Deep-redact a sample record, truncating long strings."""
    if isinstance(value, str):
        masked, hits = mask_secrets(value)
        if len(masked) > MAX_SAMPLE_CHARS:
            masked = masked[:MAX_SAMPLE_CHARS] + f"...<truncated {len(masked)} chars>"
        return masked, hits
    if isinstance(value, list):
        out, hits = [], []
        for item in value[:5]:
            m, h = redact_preview(item)
            out.append(m)
            hits.extend(h)
        if len(value) > 5:
            out.append(f"...<{len(value) - 5} more items>")
        return out, sorted(set(hits))
    if isinstance(value, dict):
        out, hits = {}, []
        for index, (key, item) in enumerate(value.items()):
            if index >= 20:
                out["..."] = f"<{len(value) - 20} more keys>"
                break
            m, h = redact_preview(item)
            out[str(key)] = m
            hits.extend(h)
        return out, sorted(set(hits))
    if value is None or isinstance(value, (int, float, bool)):
        return value, []
    return str(value), []


def merge_schema(observed: dict[str, set[str]], record: dict[str, Any]) -> None:
    for key, value in record.items():
        if isinstance(value, bool):
            kind = "bool"
        elif isinstance(value, int):
            kind = "int"
        elif isinstance(value, float):
            kind = "float"
        elif isinstance(value, list):
            kind = "list"
        elif isinstance(value, dict):
            kind = "object"
        elif isinstance(value, str):
            kind = "str"
        else:
            kind = type(value).__name__
        observed.setdefault(str(key), set()).add(kind)


def iter_jsonl(path: Path) -> Iterator[tuple[int, dict[str, Any] | None, str]]:
    """Yield (line_number, parsed_or_None, raw_line) streaming, never loading all."""
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        for line_number, raw in enumerate(handle, start=1):
            stripped = raw.strip()
            if not stripped:
                yield line_number, None, raw
                continue
            try:
                parsed = json.loads(stripped)
            except json.JSONDecodeError:
                yield line_number, None, raw
                continue
            yield line_number, parsed if isinstance(parsed, dict) else {"__value__": parsed}, raw


@dataclass
class FileReport:
    relative_path: str
    size_bytes: int
    extension: str
    kind: str
    sha256: str = ""
    record_count: int | None = None
    malformed_records: int | None = None
    schema: dict[str, list[str]] = field(default_factory=dict)
    columns: list[str] = field(default_factory=list)
    samples: list[Any] = field(default_factory=list)
    secret_flags: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "relative_path": self.relative_path,
            "size_bytes": self.size_bytes,
            "extension": self.extension,
            "kind": self.kind,
            "sha256": self.sha256,
            "record_count": self.record_count,
            "malformed_records": self.malformed_records,
            "schema": {k: sorted(v) for k, v in sorted(self.schema.items())},
            "columns": self.columns,
            "samples": self.samples,
            "secret_flags": self.secret_flags,
            "notes": self.notes,
        }


def classify(ext: str) -> str:
    if ext in TABULAR_EXT:
        return "tabular"
    if ext in ARCHIVE_EXT:
        return "archive"
    if ext in BINARY_DOC_EXT:
        return "binary_document"
    if ext in {".md", ".txt", ".rst"}:
        return "text"
    return "other"


def probe_jsonl(path: Path, report: FileReport) -> None:
    count = 0
    malformed = 0
    observed: dict[str, set[str]] = {}
    samples: list[Any] = []
    flags: set[str] = set()
    for _line_no, parsed, raw in iter_jsonl(path):
        if parsed is None:
            if raw.strip():
                malformed += 1
            continue
        count += 1
        if len(observed) < 60:
            merge_schema(observed, parsed)
        if len(samples) < MAX_SAMPLES_PER_FILE:
            preview, hits = redact_preview(parsed)
            samples.append(preview)
            flags.update(hits)
        if parsed is not None and len(samples) >= MAX_SAMPLES_PER_FILE and len(observed) >= 60:
            # Keep counting, but stop accumulating to bound memory.
            pass
    report.record_count = count
    report.malformed_records = malformed
    report.schema = observed
    report.samples = samples
    report.secret_flags = sorted(flags)


def probe_json(path: Path, report: FileReport) -> None:
    try:
        with path.open("r", encoding="utf-8", errors="replace") as handle:
            data = json.load(handle)
    except (json.JSONDecodeError, OSError) as exc:
        report.notes.append(f"json_parse_error: {type(exc).__name__}")
        return

    records: list[Any]
    if isinstance(data, list):
        records = data
        report.notes.append("top_level_list")
    elif isinstance(data, dict):
        list_key = None
        for candidate in ("data", "questions", "records", "items", "rows", "examples"):
            if isinstance(data.get(candidate), list):
                list_key = candidate
                break
        if list_key:
            records = data[list_key]
            report.notes.append(f"records_under_key:{list_key}")
            report.columns = [k for k in data.keys() if k != list_key][:20]
        else:
            records = [data]
            report.notes.append("top_level_object")
    else:
        report.notes.append(f"unexpected_top_level_type:{type(data).__name__}")
        return

    report.record_count = len(records)
    observed: dict[str, set[str]] = {}
    samples: list[Any] = []
    flags: set[str] = set()
    for record in records[:2000]:
        if isinstance(record, dict):
            merge_schema(observed, record)
    for record in records[:MAX_SAMPLES_PER_FILE]:
        preview, hits = redact_preview(record)
        samples.append(preview)
        flags.update(hits)
    report.schema = observed
    report.samples = samples
    report.secret_flags = sorted(flags)


def probe_delimited(path: Path, report: FileReport, delimiter: str) -> None:
    count = 0
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        header: list[str] | None = None
        samples: list[Any] = []
        flags: set[str] = set()
        try:
            header = next(reader)
        except StopIteration:
            report.notes.append("empty_file")
            report.record_count = 0
            return
        for row in reader:
            count += 1
            if len(samples) < MAX_SAMPLES_PER_FILE:
                paired = dict(zip(header, row)) if header else {"__row__": row}
                preview, hits = redact_preview(paired)
                samples.append(preview)
                flags.update(hits)
        report.columns = header or []
        report.record_count = count
        report.samples = samples
        report.secret_flags = sorted(flags)


def probe_xlsx(path: Path, report: FileReport) -> None:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        report.notes.append("openpyxl_not_installed:cannot_count_rows")
        return
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - third-party lib raises many types
        report.notes.append(f"xlsx_open_failed:{type(exc).__name__}")
        return
    report.notes.append(f"sheets:{','.join(workbook.sheetnames)}")
    total = 0
    samples: list[Any] = []
    flags: set[str] = set()
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = None
        sheet_rows = 0
        for row in rows:
            if header is None:
                header = [str(c) if c is not None else "" for c in row]
                report.columns = report.columns or header
                continue
            sheet_rows += 1
            if len(samples) < MAX_SAMPLES_PER_FILE:
                paired = dict(zip(header, [str(c) for c in row]))
                preview, hits = redact_preview(paired)
                samples.append(preview)
                flags.update(hits)
        total += sheet_rows
        report.notes.append(f"sheet:{sheet.title}:rows:{sheet_rows}")
    workbook.close()
    report.record_count = total
    report.samples = samples
    report.secret_flags = sorted(flags)


def probe_pdf(path: Path, report: FileReport) -> None:
    try:
        import pypdf  # type: ignore
    except ImportError:
        report.notes.append("pypdf_not_installed:page_count_unavailable")
        return
    try:
        reader = pypdf.PdfReader(str(path))
        report.record_count = len(reader.pages)
        report.notes.append(f"pdf_pages:{len(reader.pages)}")
    except Exception as exc:  # noqa: BLE001
        report.notes.append(f"pdf_open_failed:{type(exc).__name__}")


def describe_file(path: Path, root: Path) -> FileReport:
    ext = path.suffix.lower()
    report = FileReport(
        relative_path=str(path.relative_to(root)).replace("\\", "/"),
        size_bytes=path.stat().st_size,
        extension=ext,
        kind=classify(ext),
    )

    if ext in ARCHIVE_EXT:
        report.notes.append("archive:not_extracted_by_inventory")

    if ext in {".jsonl", ".ndjson"}:
        probe_jsonl(path, report)
    elif ext == ".json":
        probe_json(path, report)
    elif ext == ".csv":
        probe_delimited(path, report, ",")
    elif ext == ".tsv":
        probe_delimited(path, report, "\t")
    elif ext in {".xlsx", ".xls"}:
        probe_xlsx(path, report)
    elif ext == ".pdf":
        probe_pdf(path, report)

    report.sha256 = sha256_of(path)
    return report


def inventory(datasets_root: Path) -> dict[str, Any]:
    root = datasets_root.resolve()
    if not root.is_dir():
        raise SystemExit(f"datasets root not found: {root}")

    groups: dict[str, list[FileReport]] = {}
    ungrouped: list[FileReport] = []

    for entry in sorted(root.iterdir(), key=lambda p: p.name.lower()):
        if entry.is_dir():
            files = [
                describe_file(f, root)
                for f in sorted(entry.rglob("*"), key=lambda p: str(p).lower())
                if f.is_file()
            ]
            groups[entry.name] = files
        elif entry.is_file():
            ungrouped.append(describe_file(entry, root))

    payload = {
        "datasets_root": str(root),
        "dataset_count": len(groups),
        "dataset_names": list(groups.keys()),
        "ungrouped_file_count": len(ungrouped),
        "datasets": [
            {
                "dataset_id": name,
                "source_path": f"datasets/{name}",
                "file_count": len(files),
                "total_size_bytes": sum(f.size_bytes for f in files),
                "files": [f.to_dict() for f in files],
            }
            for name, files in groups.items()
        ],
        "ungrouped_files": [f.to_dict() for f in ungrouped],
    }
    return payload


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--datasets-root",
        default=os.environ.get("LEARNXSK_DATASETS_ROOT", str(_common.DEFAULT_DATASETS_ROOT)),
        help="Read-only datasets directory to inventory.",
    )
    parser.add_argument(
        "--out",
        default="reports/data_inventory.json",
        help="Output path relative to the training root.",
    )
    _common.add_common_args(parser)
    args = parser.parse_args(argv)

    datasets_root = Path(args.datasets_root)
    _common.info(f"inventorying (read-only): {datasets_root}")
    payload = inventory(datasets_root)

    out_path = _common.artifact_path(args.out, args.out_root)
    with out_path.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2, ensure_ascii=False)

    _common.info(f"dataset groupings discovered: {payload['dataset_count']}")
    for dataset in payload["datasets"]:
        total = dataset["total_size_bytes"]
        _common.info(
            f"  - {dataset['dataset_id']}: {dataset['file_count']} file(s), "
            f"{total / 1024 / 1024:.1f} MiB"
        )
    _common.info(f"ungrouped loose files: {payload['ungrouped_file_count']}")
    _common.info(f"wrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
